"""API tests for resume ingestion."""

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from careergraph.main import app


client = TestClient(app)


def test_extract_txt_resume() -> None:
    """The API extracts text from a TXT resume."""
    response = client.post(
        "/api/v1/resumes/extract",
        files={
            "file": (
                "resume.txt",
                b"Shruti Sharma\nPython FastAPI SQL",
                "text/plain",
            )
        },
    )

    assert response.status_code == 200

    body = response.json()

    assert body["filename"] == "resume.txt"
    assert body["media_type"] == "text/plain"
    assert body["text"] == "Shruti Sharma\nPython FastAPI SQL"
    assert body["file_size_bytes"] > 0


def test_extract_docx_resume() -> None:
    """The API extracts text from a DOCX resume."""
    document = Document()
    document.add_paragraph("Shruti Sharma")
    document.add_paragraph("Python | FastAPI | SQL")

    buffer = BytesIO()
    document.save(buffer)

    response = client.post(
        "/api/v1/resumes/extract",
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200

    body = response.json()

    assert body["filename"] == "resume.docx"
    assert "Shruti Sharma" in body["text"]
    assert "Python | FastAPI | SQL" in body["text"]


def test_reject_unsupported_resume_format() -> None:
    """The API rejects unsupported resume formats."""
    response = client.post(
        "/api/v1/resumes/extract",
        files={
            "file": (
                "resume.exe",
                b"resume",
                "application/octet-stream",
            )
        },
    )

    assert response.status_code == 415
    assert "Unsupported resume format" in response.json()["detail"]


def test_reject_empty_resume() -> None:
    """The API rejects empty resume files."""
    response = client.post(
        "/api/v1/resumes/extract",
        files={
            "file": (
                "resume.txt",
                b"",
                "text/plain",
            )
        },
    )

    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_reject_missing_resume_file() -> None:
    """The API requires a resume file."""
    response = client.post("/api/v1/resumes/extract")

    assert response.status_code == 422


def test_reject_pdf_without_extractable_text() -> None:
    """The API rejects a PDF with no extractable text."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    buffer = BytesIO()
    writer.write(buffer)

    response = client.post(
        "/api/v1/resumes/extract",
        files={
            "file": (
                "resume.pdf",
                buffer.getvalue(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 400
    assert "no extractable text" in response.json()["detail"].lower()
